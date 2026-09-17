"""Build the pptx/docx/xlsx test files Vault Scope is checked against.

Run it whenever the fixtures need regenerating:
    python fixtures/make_fixtures.py
Files land in fixtures/out/ and are deliberately outside the vault.
"""

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.util import Inches, Pt

OUT = Path(__file__).parent / "out"
OUT.mkdir(exist_ok=True)

# Text we can grep for later to prove extraction and anchoring worked.
NEEDLE = "the CIA triad is not a checklist"


def make_pptx() -> Path:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    title_layout = prs.slide_layouts[0]
    bullet_layout = prs.slide_layouts[1]

    # 1: title slide
    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = "Vault Scope Fixture Deck"
    s.placeholders[1].text = "Five slides, speaker notes, a table"
    s.notes_slide.notes_text_frame.text = "Opening notes: say hello, set the agenda."

    # 2: plain bullets
    s = prs.slides.add_slide(bullet_layout)
    s.shapes.title.text = "Bullets"
    body = s.placeholders[1].text_frame
    body.text = "First bullet"
    for line in ("Second bullet", "Third bullet with more words in it"):
        body.add_paragraph().text = line
    s.notes_slide.notes_text_frame.text = "Notes for the bullet slide."

    # 3: a table
    s = prs.slides.add_slide(blank)
    box = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    box.text_frame.text = "A table"
    rows, cols = 3, 3
    table = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.4), Inches(9), Inches(2)).table
    headers = ["Control", "Owner", "Status"]
    data = [["MFA", "IT", "Done"], ["Backups", "Ops", "Partial"]]
    for c, head in enumerate(headers):
        table.cell(0, c).text = head
    for r, row in enumerate(data, start=1):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    # 4: the slide carrying the anchor needle
    s = prs.slides.add_slide(blank)
    box = s.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(8.8), Inches(3))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = "Remember " + NEEDLE + ". Instead it is a way of asking three questions."
    tf.paragraphs[0].runs[0].font.size = Pt(28)
    s.notes_slide.notes_text_frame.text = "Slide 4 notes: this is the line to comment on."

    # 5: no notes at all, to prove the notes panel handles absence
    s = prs.slides.add_slide(blank)
    box = s.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(8.8), Inches(1.5))
    box.text_frame.text = "Last slide, no speaker notes here."

    path = OUT / "fixture-deck.pptx"
    prs.save(path)
    return path


def make_docx() -> Path:
    doc = Document()
    doc.add_heading("Vault Scope Fixture Document", level=1)
    doc.add_paragraph("A short paragraph on the first page so there is text to select.")
    doc.add_paragraph("Remember " + NEEDLE + ". Instead it is a way of asking three questions.")

    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Owner"
    table.cell(1, 1).text = "Ops"

    doc.add_page_break()
    doc.add_heading("Second page", level=2)
    for i in range(1, 6):
        doc.add_paragraph(f"Body paragraph {i} on the second page.")

    path = OUT / "fixture-doc.docx"
    doc.save(path)
    return path


def make_xlsx() -> Path:
    wb = Workbook()

    first = wb.active
    first.title = "Controls"
    first.append(["Control", "Owner", "Status"])
    for row in (["MFA", "IT", "Done"], ["Backups", "Ops", "Partial"], ["Logging", "SecOps", "Open"]):
        first.append(row)
    # A merge, so the rowspan/colspan path gets exercised.
    first.merge_cells("A6:C6")
    first["A6"] = "Merged footer row"

    second = wb.create_sheet("Numbers")
    for r in range(1, 21):
        for c in range(1, 6):
            second.cell(row=r, column=c, value=r * c)
    second["G1"] = "Column G exists to widen the used range"

    third = wb.create_sheet("Empty")
    third["A1"] = "Only one cell"

    path = OUT / "fixture-book.xlsx"
    wb.save(path)
    return path


def make_code() -> Path:
    path = OUT / "fixture-code.py"
    path.write_text(
        '"""A multi-line docstring so the highlighter has to span lines.\n'
        "Second line of the docstring.\n"
        '"""\n\n'
        "import os\n\n\n"
        "class Example:\n"
        "    def __init__(self, name: str) -> None:\n"
        "        self.name = name  # trailing comment\n\n"
        "    def greet(self) -> str:\n"
        '        return f"hello {self.name}"\n\n\n'
        'if __name__ == "__main__":\n'
        "    print(Example(os.getcwd()).greet())\n",
        encoding="utf-8",
    )
    return path


if __name__ == "__main__":
    for made in (make_pptx(), make_docx(), make_xlsx(), make_code()):
        print(f"wrote {made}")
